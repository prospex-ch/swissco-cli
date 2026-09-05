"""FINMA's authorised banks and securities firms, joined to a UID.

Two files, both published by FINMA, both static downloads, neither needing a
credential:

``beh.xlsx``
    Every bank and securities firm FINMA authorises, with its licence type,
    supervisory category and four status flags. ``?sc_lang=en`` is not
    cosmetic -- it fixes the column headers this parser matches on.

``uid.csv``
    A separate ``(name, city, authorisation type) -> UID`` crosswalk. The
    workbook itself carries no UID; without this file, nothing here joins to a
    company.

**Banks and securities firms only.** FINMA publishes two dozen further lists --
insurers, portfolio managers, fund management companies, market
infrastructures, self-regulatory organisations -- and none of them is here. The
reason is not effort. This workbook's parser finds its columns by *header name*
and checks itself against FINMA's own declared total, so a layout change stops
it loudly. The parsers for the other lists read columns by position, where the
same change silently shifts a city into the licence column and reports success.
One dataset that can tell you when it is wrong is worth more than twenty-five
that cannot.

So: a company absent from this list is not "unlicensed". It is not a bank or a
securities firm, which is a much smaller claim, and it is the only one the CLI
is entitled to make.

``openpyxl`` is imported inside :func:`parse_bank_workbook` rather than at
module scope, and :mod:`swissco.cli` imports this module inside its handlers.
``uvx swissco lookup`` must not pay for a spreadsheet library it is not going
to use.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, replace
from datetime import timedelta
from pathlib import Path
from typing import Callable

from . import _cache
from ._http import HttpClient

BASE = "https://www.finma.ch"
BANKS_URL = (
    "https://www.finma.ch/en/~/media/finma/dokumente/"
    "bewilligungstraeger/xlsx/beh.xlsx?sc_lang=en"
)
UID_URL = (
    "https://www.finma.ch/en/~/media/finma/dokumente/"
    "bewilligungstraeger/csv/uid.csv?sc_lang=en"
)

#: FINMA publishes no rate limit. One request a second is the floor the Prospex
#: collectors use against this host and it is carried over unchanged.
MIN_INTERVAL = 1.0

LICENCE_TYPES = frozenset(
    {
        "Bank",
        "Securities firm",
        "Foreign bank branch office",
        "Foreign securities firm branch office",
    }
)

CATEGORIES = frozenset({"1", "2", "3", "4", "5"})

EXPECTED_HEADERS = (
    "Name",
    "City",
    "Licensing",
    "foreign control",
    "no activity as securities firm",
    "non-account-holding securities firms",
    "about to cease operations",
    "Category",
)

# FINMA appends a parenthetical breakdown after the count, so match the count
# and let the rest of the line vary.
_TOTAL_RE = re.compile(r"^Total authorised banks and securities firms:\s*(\d+)\b.*$")


class FinmaError(Exception):
    """FINMA's published file no longer matches the shape this parser needs.

    Deliberately not a :class:`ValueError`. The CLI maps ``ValueError`` to
    ``invalid_argument``, and "FINMA reformatted their spreadsheet" is not
    something the user typed wrong.
    """


@dataclass(frozen=True)
class BankRecord:
    """One authorised institution, as the workbook lists it."""

    name: str
    city: str
    licence_type: str
    foreign_control: bool
    no_securities_firm_activity: bool
    non_account_holding_securities_firm: bool
    about_to_cease_operations: bool
    supervisory_category: str
    source_row: int
    uid: str = ""


@dataclass(frozen=True)
class BanksSnapshot:
    """Every record in one published workbook, plus the total FINMA declared."""

    records: tuple[BankRecord, ...]
    declared_total: int

    @property
    def with_uid(self) -> int:
        return sum(1 for record in self.records if record.uid)


class FinmaClient(HttpClient):
    """Downloads the two files. Nothing else on this host is ever requested."""

    allowed_hosts = frozenset({"www.finma.ch"})

    def fetch_banks(self) -> bytes:
        return self.get(
            BANKS_URL,
            accept="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ).content

    def fetch_uids(self) -> bytes:
        return self.get(UID_URL, accept="text/csv").content


# -- parsing ----------------------------------------------------------------


def normalize_source_text(value: object) -> str:
    """Collapse presentation-only differences, keeping the source's wording."""
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def crosswalk_key(name: str, city: str, licence_type: str) -> tuple[str, str, str]:
    """The three fields the crosswalk joins on, normalised and case-folded."""
    return tuple(
        normalize_source_text(value).casefold() for value in (name, city, licence_type)
    )


def parse_bank_workbook(content: bytes) -> BanksSnapshot:
    """Parse and validate one complete workbook.

    Every check here is a tripwire rather than a nicety: an unknown licence
    value, a category outside 1-5, a flag that is neither empty nor ``X``, a
    header that gained or lost a column, or a record count that disagrees with
    FINMA's own footer total all stop the parse. The alternative -- printing
    whatever survived -- is how a spreadsheet reformat becomes wrong data with
    a zero exit code.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise FinmaError(f"cannot open the FINMA workbook: {exc}") from exc

    try:
        if len(workbook.sheetnames) != 1:
            raise FinmaError(f"expected one FINMA worksheet, found {workbook.sheetnames!r}")
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_offset, positions = _find_header(rows)
    records: list[BankRecord] = []
    declared_total = None
    for offset, row in enumerate(rows[header_offset + 1 :], start=header_offset + 2):
        name = normalize_source_text(row[positions["Name"]])
        total = _TOTAL_RE.fullmatch(name)
        if total:
            declared_total = int(total.group(1))
            continue
        if not name:
            continue

        city = normalize_source_text(row[positions["City"]])
        licence_type = normalize_source_text(row[positions["Licensing"]])
        category = normalize_source_text(row[positions["Category"]])
        if not city:
            raise FinmaError(f"missing city in workbook row {offset}")
        if licence_type not in LICENCE_TYPES:
            raise FinmaError(f"unexpected licensing value {licence_type!r} in row {offset}")
        if category not in CATEGORIES:
            raise FinmaError(f"unexpected supervisory category {category!r} in row {offset}")

        records.append(
            BankRecord(
                name=name,
                city=city,
                licence_type=licence_type,
                foreign_control=_flag(row[positions["foreign control"]], offset, "foreign control"),
                no_securities_firm_activity=_flag(
                    row[positions["no activity as securities firm"]],
                    offset,
                    "no activity as securities firm",
                ),
                non_account_holding_securities_firm=_flag(
                    row[positions["non-account-holding securities firms"]],
                    offset,
                    "non-account-holding securities firms",
                ),
                about_to_cease_operations=_flag(
                    row[positions["about to cease operations"]],
                    offset,
                    "about to cease operations",
                ),
                supervisory_category=category,
                source_row=offset,
            )
        )

    if declared_total is None:
        raise FinmaError("the FINMA workbook has no declared total")
    if len(records) != declared_total:
        raise FinmaError(
            f"the FINMA workbook declares {declared_total} records but {len(records)} parsed"
        )
    return BanksSnapshot(tuple(records), declared_total)


def parse_uid_crosswalk(content: bytes) -> dict[tuple[str, str, str], str]:
    """The ``(name, city, licence type) -> UID`` mapping, exactly as published.

    Rows without a UID are dropped rather than kept as an empty string. FINMA
    publishes some authorised entities with no UID at all, and those rows are
    absence of join evidence, not evidence of absence.
    """
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FinmaError("the FINMA UID crosswalk is not UTF-8") from exc

    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=";")
    expected = {
        "Name",
        "City",
        "AuthorisationTypeDE",
        "AuthorisationTypeFR",
        "AuthorisationTypeIT",
        "AuthorisationTypeEN",
        "UID",
    }
    if set(reader.fieldnames or ()) != expected:
        raise FinmaError(f"unexpected FINMA UID columns: {reader.fieldnames!r}")

    result: dict[tuple[str, str, str], str] = {}
    for row_number, row in enumerate(reader, start=2):
        name = normalize_source_text(row["Name"])
        city = normalize_source_text(row["City"])
        licence_type = normalize_source_text(row["AuthorisationTypeEN"])
        uid = normalize_source_text(row["UID"])
        if not name or not city or not licence_type:
            raise FinmaError(f"incomplete FINMA UID row {row_number}")
        if not uid:
            continue
        key = crosswalk_key(name, city, licence_type)
        previous = result.get(key)
        if previous is not None and previous != uid:
            raise FinmaError(f"conflicting UIDs for {name!r}, {city!r}, {licence_type!r}")
        result[key] = uid
    return result


def join_uids(
    snapshot: BanksSnapshot, crosswalk: dict[tuple[str, str, str], str]
) -> BanksSnapshot:
    """*snapshot* with each record's ``uid`` filled in from *crosswalk*."""
    records = tuple(
        replace(
            record,
            uid=crosswalk.get(crosswalk_key(record.name, record.city, record.licence_type), ""),
        )
        for record in snapshot.records
    )
    return replace(snapshot, records=records)


def _flag(value: object, row: int, column: str) -> bool:
    text = normalize_source_text(value)
    if text not in {"", "X"}:
        raise FinmaError(f"unexpected {column} flag {text!r} in workbook row {row}")
    return text == "X"


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for offset, row in enumerate(rows):
        values = [normalize_source_text(value) for value in row]
        if values[:3] != ["Name", "City", "Licensing"]:
            continue
        positions = {value: index for index, value in enumerate(values) if value}
        missing = [header for header in EXPECTED_HEADERS if header not in positions]
        if missing:
            raise FinmaError(f"the FINMA workbook header is missing: {', '.join(missing)}")
        unexpected = [header for header in positions if header not in EXPECTED_HEADERS]
        if unexpected:
            raise FinmaError(
                f"the FINMA workbook header has unexpected columns: {', '.join(unexpected)}"
            )
        return offset, positions
    raise FinmaError("the FINMA workbook has no Name/City/Licensing header row")


# -- fetching, with the cache in front ---------------------------------------


def snapshot(
    client: FinmaClient,
    *,
    cache_dir: Path | None = None,
    max_age: timedelta = _cache.DEFAULT_MAX_AGE,
    use_cache: bool = True,
    on_note: Callable[[str], None] | None = None,
) -> BanksSnapshot:
    """The joined snapshot, from the cache when it is fresh enough.

    Each file is parsed before it is written back to the cache, so a FINMA
    layout change leaves the last good copy on disk instead of replacing it
    with something that will not parse tomorrow either.
    """
    workbook = _file(
        BANKS_URL, "beh.xlsx", client.fetch_banks, cache_dir, max_age, use_cache, on_note
    )
    banks = parse_bank_workbook(workbook.content)
    workbook.keep()

    crosswalk_file = _file(
        UID_URL, "uid.csv", client.fetch_uids, cache_dir, max_age, use_cache, on_note
    )
    crosswalk = parse_uid_crosswalk(crosswalk_file.content)
    crosswalk_file.keep()

    return join_uids(banks, crosswalk)


@dataclass
class _Fetched:
    content: bytes
    path: Path | None
    url: str
    from_cache: bool

    def keep(self) -> None:
        """Write the bytes to the cache, now that they are known to parse."""
        if self.path is not None and not self.from_cache:
            _cache.write(self.path, self.content, url=self.url)


def _file(url, name, fetch, cache_dir, max_age, use_cache, on_note) -> _Fetched:
    if cache_dir is None:
        return _Fetched(fetch(), None, url, False)
    path = cache_dir / "finma" / name
    before = _cache.read(path, max_age=max_age) if use_cache else None
    content = _cache.download(
        url,
        path,
        fetch=lambda _: fetch(),
        max_age=max_age,
        use_cache=use_cache,
        on_note=on_note,
    )
    return _Fetched(content, path, url, from_cache=before is not None)


# -- selecting and rendering -------------------------------------------------


def by_uid(snap: BanksSnapshot, uid: str) -> BankRecord | None:
    """The record carrying *uid*, or ``None``.

    Compared with the punctuation stripped: ``companies.resolve_uid`` hands
    back ``CHE105845287`` and FINMA publishes ``CHE-105.845.287``. They are the
    same UID, and a string comparison would say otherwise.
    """
    wanted = squash(uid)
    for record in snap.records:
        if squash(record.uid) == wanted:
            return record
    return None


def squash(uid: str) -> str:
    """A UID with its separators removed, for comparing two spellings of one."""
    return "".join(character for character in (uid or "").upper() if character.isalnum())


def filter_records(
    snap: BanksSnapshot,
    *,
    query: str = "",
    licence_type: str = "",
    category: str = "",
) -> list[BankRecord]:
    """Records matching a case-insensitive name/city substring and the filters."""
    needle = normalize_source_text(query).casefold()
    out = []
    for record in snap.records:
        if needle and needle not in record.name.casefold() and needle not in record.city.casefold():
            continue
        if licence_type and record.licence_type != licence_type:
            continue
        if category and record.supervisory_category != category:
            continue
        out.append(record)
    return out


def bank_row(record: BankRecord) -> dict:
    """One record as a row. Key order is column order."""
    return {
        "name": record.name,
        "city": record.city,
        "licence_type": record.licence_type,
        "supervisory_category": record.supervisory_category,
        "uid": record.uid,
        "foreign_control": record.foreign_control,
        "no_securities_firm_activity": record.no_securities_firm_activity,
        "non_account_holding_securities_firm": record.non_account_holding_securities_firm,
        "about_to_cease_operations": record.about_to_cease_operations,
    }


def lookup_fields(record: BankRecord) -> dict:
    """The subset ``swissco lookup --finma`` adds to a company row."""
    flags = [
        label
        for label, on in (
            ("foreign control", record.foreign_control),
            ("no securities-firm activity", record.no_securities_firm_activity),
            ("non-account-holding", record.non_account_holding_securities_firm),
            ("about to cease operations", record.about_to_cease_operations),
        )
        if on
    ]
    return {
        "finma_licence": record.licence_type,
        "finma_category": record.supervisory_category,
        "finma_city": record.city,
        "finma_flags": flags,
    }
