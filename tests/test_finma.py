"""The FINMA workbook parser, its crosswalk, and the join between them.

The success paths run against the two files as FINMA publishes them. The
failure paths cannot: FINMA does not serve a workbook with a missing column, so
those cases are built in memory. Every one of them exists to prove the parser
*stops* rather than printing whatever survived, which is the only reason this
dataset was chosen over FINMA's two dozen others.
"""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import Workbook

from swissco import finma


def workbook(rows, *, declared_total=None, headers=None, sheets=1) -> bytes:
    """A workbook in FINMA's shape, for the failures FINMA does not publish."""
    book = Workbook()
    sheet = book.active
    sheet.append(list(headers or finma.EXPECTED_HEADERS))
    for row in rows:
        sheet.append(list(row))
    total = len(rows) if declared_total is None else declared_total
    sheet.append([f"Total authorised banks and securities firms: {total}"])
    for extra in range(sheets - 1):
        book.create_sheet(f"extra{extra}")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def crosswalk(rows) -> bytes:
    """A crosswalk CSV in FINMA's shape: semicolons, UTF-8 with a BOM."""
    fields = [
        "Name",
        "City",
        "AuthorisationTypeDE",
        "AuthorisationTypeFR",
        "AuthorisationTypeIT",
        "AuthorisationTypeEN",
        "UID",
    ]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fields, delimiter=";", lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({field: row.get(field, "") for field in fields})
    return buffer.getvalue().encode("utf-8-sig")


BANK = ("Musterbank AG", "Bern", "Bank", "", "", "", "", "3")


class TestTheWorkbookAsPublished:
    def test_the_record_count_matches_finmas_own_total(self, finma_workbook):
        snapshot = finma.parse_bank_workbook(finma_workbook)
        assert len(snapshot.records) == snapshot.declared_total
        assert snapshot.declared_total > 0

    def test_every_licence_and_category_is_one_finma_documents(self, finma_workbook):
        snapshot = finma.parse_bank_workbook(finma_workbook)
        assert {r.licence_type for r in snapshot.records} <= finma.LICENCE_TYPES
        assert {r.supervisory_category for r in snapshot.records} <= finma.CATEGORIES

    def test_every_record_has_a_name_and_a_city(self, finma_workbook):
        snapshot = finma.parse_bank_workbook(finma_workbook)
        assert all(record.name and record.city for record in snapshot.records)


class TestTheCrosswalkAsPublished:
    def test_it_maps_the_three_key_fields_to_a_uid(self, finma_crosswalk):
        mapping = finma.parse_uid_crosswalk(finma_crosswalk)
        assert mapping
        key, uid = next(iter(mapping.items()))
        assert len(key) == 3
        assert uid.startswith("CHE")

    def test_nearly_every_bank_joins(self, finma_workbook, finma_crosswalk):
        joined = finma.join_uids(
            finma.parse_bank_workbook(finma_workbook),
            finma.parse_uid_crosswalk(finma_crosswalk),
        )
        # FINMA publishes a handful of authorised entities with no UID at all.
        # A join rate this high is what makes the UID lookup worth offering; if
        # it collapses, the crosswalk's key fields have drifted.
        assert joined.with_uid / len(joined.records) > 0.9

    def test_a_uid_is_found_whichever_way_it_is_punctuated(
        self, finma_workbook, finma_crosswalk
    ):
        joined = finma.join_uids(
            finma.parse_bank_workbook(finma_workbook),
            finma.parse_uid_crosswalk(finma_crosswalk),
        )
        record = next(r for r in joined.records if r.uid)
        assert finma.by_uid(joined, record.uid) is record
        assert finma.by_uid(joined, finma.squash(record.uid)) is record
        assert finma.by_uid(joined, record.uid.lower()) is record


class TestTheWorkbookTripwires:
    def test_a_missing_column_stops_the_parse(self):
        headers = [h for h in finma.EXPECTED_HEADERS if h != "Category"]
        with pytest.raises(finma.FinmaError, match="missing"):
            finma.parse_bank_workbook(workbook([BANK[:-1]], headers=headers))

    def test_an_extra_column_stops_the_parse(self):
        headers = (*finma.EXPECTED_HEADERS, "Something new")
        with pytest.raises(finma.FinmaError, match="unexpected columns"):
            finma.parse_bank_workbook(workbook([(*BANK, "x")], headers=headers))

    def test_a_second_worksheet_stops_the_parse(self):
        with pytest.raises(finma.FinmaError, match="one FINMA worksheet"):
            finma.parse_bank_workbook(workbook([BANK], sheets=2))

    def test_no_header_row_stops_the_parse(self):
        with pytest.raises(finma.FinmaError, match="no Name/City/Licensing"):
            finma.parse_bank_workbook(workbook([BANK], headers=["a", "b", "c"] + [""] * 5))

    def test_an_unknown_licence_value_stops_the_parse(self):
        row = ("Musterbank AG", "Bern", "Crypto bank", "", "", "", "", "3")
        with pytest.raises(finma.FinmaError, match="licensing value"):
            finma.parse_bank_workbook(workbook([row]))

    def test_a_category_outside_one_to_five_stops_the_parse(self):
        row = ("Musterbank AG", "Bern", "Bank", "", "", "", "", "6")
        with pytest.raises(finma.FinmaError, match="supervisory category"):
            finma.parse_bank_workbook(workbook([row]))

    def test_a_flag_that_is_not_x_stops_the_parse(self):
        row = ("Musterbank AG", "Bern", "Bank", "Y", "", "", "", "3")
        with pytest.raises(finma.FinmaError, match="foreign control"):
            finma.parse_bank_workbook(workbook([row]))

    def test_a_missing_city_stops_the_parse(self):
        row = ("Musterbank AG", "", "Bank", "", "", "", "", "3")
        with pytest.raises(finma.FinmaError, match="missing city"):
            finma.parse_bank_workbook(workbook([row]))

    def test_a_total_that_disagrees_stops_the_parse(self):
        with pytest.raises(finma.FinmaError, match="declares 7"):
            finma.parse_bank_workbook(workbook([BANK], declared_total=7))

    def test_no_declared_total_stops_the_parse(self):
        book = Workbook()
        book.active.append(list(finma.EXPECTED_HEADERS))
        book.active.append(list(BANK))
        buffer = io.BytesIO()
        book.save(buffer)
        with pytest.raises(finma.FinmaError, match="no declared total"):
            finma.parse_bank_workbook(buffer.getvalue())

    def test_something_that_is_not_a_workbook_stops_the_parse(self):
        with pytest.raises(finma.FinmaError, match="cannot open"):
            finma.parse_bank_workbook(b"not a spreadsheet")

    def test_a_flag_reads_as_a_bool(self):
        row = ("Musterbank AG", "Bern", "Bank", "X", "", "X", "", "3")
        record = finma.parse_bank_workbook(workbook([row])).records[0]
        assert record.foreign_control is True
        assert record.no_securities_firm_activity is False
        assert record.non_account_holding_securities_firm is True


class TestTheCrosswalkTripwires:
    def test_conflicting_uids_for_one_key_stop_the_parse(self):
        rows = [
            {"Name": "A", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-1"},
            {"Name": "A", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-2"},
        ]
        with pytest.raises(finma.FinmaError, match="Conflicting|conflicting"):
            finma.parse_uid_crosswalk(crosswalk(rows))

    def test_the_same_uid_twice_is_not_a_conflict(self):
        rows = [
            {"Name": "A", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-1"},
            {"Name": "A", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-1"},
        ]
        assert len(finma.parse_uid_crosswalk(crosswalk(rows))) == 1

    def test_a_row_without_a_uid_is_dropped_rather_than_kept_empty(self):
        rows = [
            {"Name": "A", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": ""},
            {"Name": "B", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-2"},
        ]
        mapping = finma.parse_uid_crosswalk(crosswalk(rows))
        assert list(mapping.values()) == ["CHE-2"]

    def test_an_incomplete_row_stops_the_parse(self):
        rows = [{"Name": "", "City": "Bern", "AuthorisationTypeEN": "Bank", "UID": "CHE-1"}]
        with pytest.raises(finma.FinmaError, match="incomplete"):
            finma.parse_uid_crosswalk(crosswalk(rows))

    def test_unexpected_columns_stop_the_parse(self):
        content = b"Name;City;UID\nA;Bern;CHE-1\n"
        with pytest.raises(finma.FinmaError, match="unexpected FINMA UID columns"):
            finma.parse_uid_crosswalk(content)

    def test_a_file_that_is_not_utf8_stops_the_parse(self):
        with pytest.raises(finma.FinmaError, match="not UTF-8"):
            finma.parse_uid_crosswalk(b"\xff\xfe\x00Name;City")


class TestTheJoin:
    def test_the_key_ignores_case_and_spacing(self):
        assert finma.crosswalk_key("Bank  AG", "Bern", "Bank") == finma.crosswalk_key(
            "bank ag", "BERN", "bank"
        )

    def test_a_bank_with_no_crosswalk_row_keeps_an_empty_uid(self):
        snapshot = finma.parse_bank_workbook(workbook([BANK]))
        joined = finma.join_uids(snapshot, {})
        assert joined.records[0].uid == ""
        assert joined.with_uid == 0

    def test_the_uid_comes_from_the_matching_key(self):
        snapshot = finma.parse_bank_workbook(workbook([BANK]))
        mapping = {finma.crosswalk_key("Musterbank AG", "Bern", "Bank"): "CHE-123"}
        assert finma.join_uids(snapshot, mapping).records[0].uid == "CHE-123"


class TestSelecting:
    def test_a_query_matches_the_name_or_the_city(self):
        rows = [BANK, ("Zweitbank AG", "Zug", "Bank", "", "", "", "", "4")]
        snapshot = finma.parse_bank_workbook(workbook(rows))
        assert len(finma.filter_records(snapshot, query="muster")) == 1
        assert len(finma.filter_records(snapshot, query="zug")) == 1
        assert len(finma.filter_records(snapshot, query="bank")) == 2

    def test_licence_and_category_narrow_further(self):
        rows = [BANK, ("Zweitbank AG", "Zug", "Securities firm", "", "", "", "", "4")]
        snapshot = finma.parse_bank_workbook(workbook(rows))
        assert len(finma.filter_records(snapshot, licence_type="Bank")) == 1
        assert len(finma.filter_records(snapshot, category="4")) == 1

    def test_a_row_keeps_its_keys_in_column_order(self):
        record = finma.parse_bank_workbook(workbook([BANK])).records[0]
        assert list(finma.bank_row(record))[:5] == [
            "name",
            "city",
            "licence_type",
            "supervisory_category",
            "uid",
        ]

    def test_lookup_fields_name_the_flags_that_are_set(self):
        row = ("Musterbank AG", "Bern", "Bank", "X", "", "", "X", "3")
        record = finma.parse_bank_workbook(workbook([row])).records[0]
        fields = finma.lookup_fields(record)
        assert fields["finma_licence"] == "Bank"
        assert fields["finma_flags"] == ["foreign control", "about to cease operations"]
